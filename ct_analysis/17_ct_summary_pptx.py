"""One short PowerPoint about the SCAPIS CT work: data, selection, segmentation, results.

Everything is read from the files that the earlier scripts already wrote, so no
image is opened again:

    ct_3mm_noncontrast.xlsx / af_cohort.xlsx   Summary and Per_Site sheets (scripts 06, 09)
    SCAPIS clinical CSV                        AF, heart failure and stroke labels
    NIfTI folder                               how many volumes exist (file count)
    radiomic_features.xlsx                     radiomics per patient (script 14)
    fat_la_features.xlsx                       fat HU measurements (script 11)
    dice_dinov2_ct.csv                         DINOv2 + LoRA test Dice (script 16), optional
    nnunet_dice.csv                            nnU-Net test Dice, optional, columns case,dice

The prediction part is deliberately simple: one model (logistic regression) and
5-fold stratified cross-validated AUC, so three tiers can be compared directly:

    1 plain measurements   volumes and mean attenuation
    2 composite indices    CMI1, CMI2, LA-EAT fraction, fat/LA ratio
    3 radiomics            best 1-3 feature model from a small pool

Output: <output>\\SCAPIS_CT_summary.pptx and <output>\\pptx_figures\\*.png

Usage
    python 17_ct_summary_pptx.py
    python 17_ct_summary_pptx.py --dino-dice "Q:\\...\\LA82\\dinov2_lora\\dice_dinov2_ct.csv"

In a Jupyter cell:
    from importlib import import_module
    import_module("17_ct_summary_pptx").main([])
"""

from __future__ import annotations

import argparse
import itertools
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from scipy.stats import mannwhitneyu
from sklearn.linear_model import LogisticRegression
from sklearn.model_selection import StratifiedKFold, cross_val_score
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler

CT = Path(r"Q:\users\marfi\CT_analysis")
SEG = Path(r"Q:\users\marfi\Automatic segmentations\finalllll\final")
ANALYSIS = SEG / "test80_analysis"

MEASUREMENTS = [
    "eat_volume_ml",
    "la_volume_ml",
    "laeat_volume_ml",
    "eat_mean",
    "laeat_mean",
    "fat_mean_hu",
]
INDICES = ["cmi1", "cmi2", "laeat_fraction_of_eat", "fat_to_la_ratio"]
NICE = {
    "eat_volume_ml": "EAT volume (ml)",
    "la_volume_ml": "LA volume (ml)",
    "laeat_volume_ml": "LA-EAT volume (ml)",
    "fat_volume_ml": "Cardiac fat volume (ml)",
    "eat_mean": "EAT mean attenuation (HU)",
    "laeat_mean": "LA-EAT mean attenuation (HU)",
    "la_mean": "LA mean attenuation (HU)",
    "fat_mean_hu": "Fat mean attenuation (HU)",
    "fat_sd_hu": "Fat attenuation SD (HU)",
    "cmi1": "CMI1 (EAT x LA)",
    "cmi2": "CMI2 (EAT^2 x LA^0.1 x LA-EAT^3)",
    "laeat_fraction_of_eat": "LA-EAT / EAT fraction",
    "fat_to_la_ratio": "Fat / LA ratio",
    "fat_index_ml_per_cm": "Fat per cm of heart (ml/cm)",
}
PARTS = {"la": "LA", "laeat": "LA-EAT", "eat": "EAT"}
STATS = {
    "mad": "MAD",
    "robust_mad": "robust MAD",
    "iqr": "IQR",
    "variance": "variance",
    "sd": "SD",
    "mean": "mean attenuation",
    "skewness": "skewness",
    "kurtosis": "kurtosis",
    "entropy": "entropy",
    "zone_entropy": "GLSZM zone entropy",
    "zone_percentage": "GLSZM zone percentage",
    "zone_size_non_uniformity": "GLSZM SZNUN",
    "grey_level_non_uniformity": "GLSZM GLNUN",
    "small_area_emphasis": "GLSZM small area emphasis",
    "large_area_emphasis": "GLSZM large area emphasis",
}


def nice(name: str) -> str:
    """Readable label for a feature column."""
    if name in NICE:
        return NICE[name]
    for part, label in PARTS.items():
        if name.startswith(part + "_"):
            rest = name[len(part) + 1 :]
            if rest in STATS:
                return f"{label} {STATS[rest]}"
            if rest == "volume_ml":
                return f"{label} volume (ml)"
    return name.replace("_", " ")


def parse_args(argv=None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--noncontrast-cohort", type=Path, default=CT / "ct_3mm_noncontrast.xlsx")
    parser.add_argument("--contrast-cohort", type=Path, default=CT / "af_cohort.xlsx")
    parser.add_argument(
        "--clinical",
        type=Path,
        default=CT / "SCAPIS1-DATA-PETITION-1110-20260407(in).csv",
    )
    parser.add_argument("--nifti-dir", type=Path, default=Path(r"W:\SCAPIS_seg_lowres_3mm_noncontrast"))
    parser.add_argument("--features", type=Path, default=ANALYSIS / "radiomic_features.xlsx")
    parser.add_argument("--extra-features", type=Path, default=ANALYSIS / "fat_la_features.xlsx")
    parser.add_argument("--dino-dice", type=Path, default=SEG / "LA82" / "dinov2_lora" / "dice_dinov2_ct.csv")
    parser.add_argument("--nnunet-dice", type=Path, default=SEG / "LA82" / "nnunet_dice.csv")
    parser.add_argument("--manual-contrast", default="all", help="text for the manual review slide")
    parser.add_argument("--manual-noncontrast", default="about 1000")
    parser.add_argument("--output", type=Path, default=ANALYSIS)
    return parser.parse_known_args(argv)[0]


# ----------------------------------------------------------------- data tables


def read_sheet(path: Path, sheet: str) -> pd.DataFrame | None:
    if not path.exists():
        print(f"missing {path}")
        return None
    try:
        return pd.read_excel(path, sheet_name=sheet)
    except Exception as error:  # noqa: BLE001 - a missing sheet must not stop the deck
        print(f"{path.name}: no sheet {sheet} ({error})")
        return None


def yes(series: pd.Series) -> pd.Series:
    return series.astype(str).str.strip().str.upper().eq("YES")


def clinical_groups(path: Path) -> pd.DataFrame | None:
    """patient_id plus AF, heart failure and stroke as booleans."""
    if not path.exists():
        print(f"missing {path}")
        return None
    raw = pd.read_csv(path, dtype=str, low_memory=False, encoding="utf-8-sig")
    column = {c.lower(): c for c in raw.columns}
    subject = column.get("subject") or next(iter(raw.columns))
    table = pd.DataFrame({"patient_id": raw[subject].astype(str).str.strip()})
    for name, key in [
        ("AF", "af_ct_baseline"),
        ("HF", "prev_hf"),
        ("stroke", "ischemic_stroke_ic_baseline"),
    ]:
        table[name] = yes(raw[column[key]]) if key in column else False
    table["control"] = ~(table["AF"] | table["HF"] | table["stroke"])
    return table


def per_site_table(cohort: Path, clinical: pd.DataFrame | None) -> pd.DataFrame | None:
    """Series, selected patients and disease counts per site."""
    per_site = read_sheet(cohort, "Per_Site")
    selected = read_sheet(cohort, "Selected")
    if per_site is None:
        return None
    per_site = per_site.copy()
    if selected is not None and clinical is not None and "site_name" in selected.columns:
        merged = selected[["patient_id", "site_number", "site_name"]].merge(
            clinical, on="patient_id", how="left"
        )
        counts = (
            merged.groupby(["site_number", "site_name"])[["AF", "HF", "stroke", "control"]]
            .sum()
            .astype(int)
            .reset_index()
        )
        per_site["site_number"] = per_site["site_number"].astype(str)
        counts["site_number"] = counts["site_number"].astype(str)
        per_site = per_site.merge(
            counts.drop(columns=["site_name"]), on="site_number", how="left"
        )
    return per_site.fillna(0)


def feature_table(args: argparse.Namespace) -> pd.DataFrame | None:
    """Radiomics plus the fat HU measurements, one row per patient, with a group."""
    table = read_sheet(args.features, "Features")
    if table is None:
        return None
    extra = read_sheet(args.extra_features, "Features")
    if extra is not None:
        new = [c for c in extra.columns if c not in table.columns or c == "patient_id"]
        table = table.merge(extra[new], on="patient_id", how="left")
    if "group" not in table.columns:
        print("no group column in the features sheet")
        return None
    table["group"] = table["group"].astype(str)
    return table


# -------------------------------------------------------------------- analysis


def auc(table: pd.DataFrame, names: list[str], group: str) -> float:
    """5-fold stratified cross-validated AUC, logistic regression, group vs control."""
    part = table[table["group"].isin([group, "control"])]
    data = part[names].apply(pd.to_numeric, errors="coerce")
    keep = data.notna().all(axis=1)
    data, target = data[keep], (part.loc[keep, "group"] == group).astype(int)
    if target.nunique() < 2 or target.sum() < 5 or (len(target) - target.sum()) < 5:
        return float("nan")
    model = make_pipeline(StandardScaler(), LogisticRegression(max_iter=2000))
    folds = StratifiedKFold(n_splits=5, shuffle=True, random_state=0)
    scores = cross_val_score(model, data, target, cv=folds, scoring="roc_auc")
    return float(np.mean(scores))


def radiomic_columns(table: pd.DataFrame) -> list[str]:
    numeric = table.select_dtypes("number").columns
    skip = set(MEASUREMENTS) | set(INDICES) | {"age", "bmi", "sex_male"}
    return [
        c
        for c in numeric
        if c not in skip
        and any(c.startswith(part + "_") for part in PARTS)
        and not c.endswith("volume_ml")
    ]


def best_radiomic_model(
    table: pd.DataFrame, group: str, pool_size: int = 8, max_size: int = 3
) -> tuple[list[str], float]:
    """Best 1-3 feature radiomic model, pool chosen by single feature AUC."""
    single = [(c, auc(table, [c], group)) for c in radiomic_columns(table)]
    single = [row for row in single if not np.isnan(row[1])]
    single.sort(key=lambda row: row[1], reverse=True)
    pool = [name for name, _ in single[:pool_size]]
    best: tuple[list[str], float] = ([], float("nan"))
    for size in range(1, max_size + 1):
        for names in itertools.combinations(pool, size):
            value = auc(table, list(names), group)
            if not np.isnan(value) and (np.isnan(best[1]) or value > best[1]):
                best = (list(names), value)
    return best


def tiers(table: pd.DataFrame, groups: list[str]) -> pd.DataFrame:
    """One row per group and tier: plain measurement, index, radiomic model."""
    rows = []
    for group in groups:
        for name in [c for c in MEASUREMENTS if c in table.columns]:
            rows.append((group, "measurement", nice(name), auc(table, [name], group)))
        for name in [c for c in INDICES if c in table.columns]:
            rows.append((group, "index", nice(name), auc(table, [name], group)))
        names, value = best_radiomic_model(table, group)
        if names:
            rows.append(
                (group, "radiomics", " + ".join(nice(n) for n in names), value)
            )
    return pd.DataFrame(rows, columns=["group", "tier", "model", "auc"]).dropna(
        subset=["auc"]
    )


def group_tests(table: pd.DataFrame, groups: list[str]) -> pd.DataFrame:
    """Mann-Whitney of each plain measurement against the controls."""
    rows = []
    control = table[table["group"] == "control"]
    for group in groups:
        part = table[table["group"] == group]
        for name in [c for c in MEASUREMENTS if c in table.columns]:
            a = pd.to_numeric(part[name], errors="coerce").dropna()
            b = pd.to_numeric(control[name], errors="coerce").dropna()
            if len(a) < 3 or len(b) < 3:
                continue
            p = mannwhitneyu(a, b, alternative="two-sided").pvalue
            rows.append((group, nice(name), a.median(), b.median(), p))
    return pd.DataFrame(
        rows, columns=["group", "feature", "median_group", "median_control", "p_value"]
    )


# --------------------------------------------------------------------- figures


def figure_sites(per_site: pd.DataFrame, path: Path) -> Path | None:
    if per_site is None or per_site.empty:
        return None
    labels = per_site["site_name"].astype(str)
    figure, axes = plt.subplots(1, 2, figsize=(11, 4))
    axes[0].bar(labels, per_site.get("patients_selected", 0), color="#3b6ea5")
    axes[0].set_title("Patients with a selected series")
    disease = [c for c in ["AF", "HF", "stroke"] if c in per_site.columns]
    bottom = np.zeros(len(per_site))
    for name, colour in zip(disease, ["#b2182b", "#ef8a62", "#67a9cf"]):
        values = per_site[name].to_numpy(dtype=float)
        axes[1].bar(labels, values, bottom=bottom, label=name, color=colour)
        bottom += values
    axes[1].set_title("Disease cases per site")
    if disease:
        axes[1].legend()
    for axis in axes:
        axis.tick_params(axis="x", rotation=30)
        axis.set_ylabel("patients")
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)
    return path


def figure_measurements(table: pd.DataFrame, groups: list[str], path: Path) -> Path | None:
    names = [c for c in MEASUREMENTS if c in table.columns][:4]
    if not names:
        return None
    order = ["control", *groups]
    figure, axes = plt.subplots(1, len(names), figsize=(4 * len(names), 4))
    axes = np.atleast_1d(axes)
    for axis, name in zip(axes, names):
        data = [
            pd.to_numeric(table.loc[table["group"] == g, name], errors="coerce").dropna()
            for g in order
        ]
        axis.boxplot(data, showfliers=False)
        axis.set_xticks(range(1, len(order) + 1))
        axis.set_xticklabels(order)
        axis.set_title(nice(name), fontsize=10)
        axis.tick_params(axis="x", rotation=20)
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)
    return path


def figure_auc(scores: pd.DataFrame, path: Path) -> Path | None:
    if scores.empty:
        return None
    groups = sorted(scores["group"].unique())
    colour = {"measurement": "#b2182b", "index": "#ef8a62", "radiomics": "#3b6ea5"}
    figure, axes = plt.subplots(1, len(groups), figsize=(5.5 * len(groups), 4.6))
    axes = np.atleast_1d(axes)
    for axis, group in zip(axes, groups):
        part = scores[scores["group"] == group].sort_values("auc")
        axis.barh(
            range(len(part)),
            part["auc"],
            color=[colour.get(t, "#999999") for t in part["tier"]],
        )
        axis.set_yticks(range(len(part)))
        axis.set_yticklabels(part["model"], fontsize=8)
        axis.axvline(0.5, color="black", linestyle=":")
        axis.set_xlim(0.4, 1.0)
        axis.set_xlabel("AUC, 5-fold CV")
        axis.set_title(f"{group} vs control")
    handles = [plt.Rectangle((0, 0), 1, 1, color=c) for c in colour.values()]
    figure.legend(handles, colour.keys(), loc="lower center", ncol=3)
    figure.tight_layout(rect=(0, 0.06, 1, 1))
    figure.savefig(path, dpi=150)
    plt.close(figure)
    return path


def figure_dice(dino: pd.DataFrame | None, nnunet: pd.DataFrame | None, path: Path):
    bars = []
    for name, frame in [("nnU-Net", nnunet), ("DINOv2 + LoRA", dino)]:
        if frame is None or frame.empty:
            continue
        column = next(
            (c for c in frame.columns if "dice" in str(c).lower() and frame[c].dtype != object),
            None,
        )
        if column is None:
            continue
        values = pd.to_numeric(frame[column], errors="coerce").dropna()
        if len(values):
            bars.append((name, values.mean(), values.std(), len(values)))
    if not bars:
        return None, bars
    figure, axis = plt.subplots(figsize=(5, 4))
    axis.bar(
        [b[0] for b in bars],
        [b[1] for b in bars],
        yerr=[b[2] for b in bars],
        capsize=5,
        color=["#3b6ea5", "#ef8a62"][: len(bars)],
    )
    axis.set_ylim(0, 1)
    axis.set_ylabel("Dice on the test set")
    for index, bar in enumerate(bars):
        axis.text(index, bar[1] + 0.03, f"{bar[1]:.3f}", ha="center")
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)
    return path, bars


# ------------------------------------------------------------------ PowerPoint


def add_title(presentation: Presentation, title: str, subtitle: str) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets(presentation: Presentation, title: str, lines: list[str]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    frame = slide.placeholders[1].text_frame
    frame.clear()
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(16)


def add_table(presentation: Presentation, title: str, frame: pd.DataFrame) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = title
    rows, columns = frame.shape
    shape = slide.shapes.add_table(
        rows + 1, columns, Inches(0.4), Inches(1.5), Inches(9.2), Inches(0.4 * (rows + 1))
    )
    table = shape.table
    for index, name in enumerate(frame.columns):
        cell = table.cell(0, index)
        cell.text = str(name).replace("_", " ")
        cell.text_frame.paragraphs[0].font.size = Pt(11)
    for row in range(rows):
        for column in range(columns):
            value = frame.iat[row, column]
            if isinstance(value, float):
                value = f"{value:,.3f}" if abs(value) < 100 else f"{value:,.0f}"
            elif isinstance(value, (int, np.integer)):
                value = f"{int(value):,}"
            cell = table.cell(row + 1, column)
            cell.text = str(value)
            cell.text_frame.paragraphs[0].font.size = Pt(11)


def add_picture(presentation: Presentation, title: str, image: Path, note: str = "") -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = title
    slide.shapes.add_picture(str(image), Inches(0.4), Inches(1.5), width=Inches(9.2))
    if note:
        box = slide.shapes.add_textbox(Inches(0.4), Inches(6.4), Inches(9.2), Inches(0.6))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = note
        paragraph.font.size = Pt(12)


# ------------------------------------------------------------------------ main


def main(argv=None) -> None:
    args = parse_args(argv)
    figures = args.output / "pptx_figures"
    figures.mkdir(parents=True, exist_ok=True)

    clinical = clinical_groups(args.clinical)
    noncontrast_summary = read_sheet(args.noncontrast_cohort, "Summary")
    contrast_summary = read_sheet(args.contrast_cohort, "Summary")
    per_site = per_site_table(args.noncontrast_cohort, clinical)
    nifti = (
        len(list(args.nifti_dir.glob("*.nii.gz"))) if args.nifti_dir.exists() else None
    )
    table = feature_table(args)

    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    add_title(
        presentation,
        "SCAPIS cardiac CT: data, segmentation and first results",
        "Six sites, automatic series selection, LA and epicardial fat, radiomics",
    )

    counts = []
    for name, summary in [("non-contrast", noncontrast_summary), ("contrast", contrast_summary)]:
        if summary is None:
            continue
        for item, value in summary.itertuples(index=False):
            counts.append((name, str(item), value))
    if counts:
        add_table(
            presentation,
            "Data and automatic series selection",
            pd.DataFrame(counts, columns=["cohort", "item", "count"]),
        )
    if nifti is not None:
        add_bullets(
            presentation,
            "NIfTI volumes generated",
            [f"{nifti:,} volumes in {args.nifti_dir}", "one volume per patient, flat folder"],
        )

    add_bullets(
        presentation,
        "How the automatic check worked",
        [
            "Every DICOM series was read once and indexed, no image was converted first.",
            "Automatic rejection rules: wrong modality, localiser or dose report, wrong",
            "   contrast state, slice thickness outside 2.0-3.5 mm, non-original",
            "   reconstruction, duplicated or missing slice positions.",
            "Coverage metric: number of unique slice positions, fewer than 40 (about 120 mm)",
            "   was dropped, so a patient is never represented by half a heart.",
            "Every dropped series is kept with its reason in the Rejected sheet, and the",
            "   poor coverage ones in Too_Few_Slices, so nothing disappears silently.",
            f"Manual review: contrast series {args.manual_contrast}, non-contrast",
            f"   {args.manual_noncontrast} inspected; the rest passed the automatic rules only.",
        ],
    )

    if per_site is not None:
        add_table(presentation, "Per site", per_site)
        image = figure_sites(per_site, figures / "sites.png")
        if image:
            add_picture(presentation, "Per site", image)

    if clinical is not None:
        add_table(
            presentation,
            "Clinical groups in the whole cohort",
            pd.DataFrame(
                [(name, int(clinical[name].sum())) for name in ["AF", "HF", "stroke", "control"]],
                columns=["group", "patients"],
            ),
        )

    dino = pd.read_csv(args.dino_dice) if args.dino_dice.exists() else None
    nnunet = pd.read_csv(args.nnunet_dice) if args.nnunet_dice.exists() else None
    dice_image, bars = figure_dice(dino, nnunet, figures / "dice.png")
    segmentation = [
        "LA and epicardial fat segmented with nnU-Net, trained on the manual labels",
        "   (82 training cases, 10 held-out test cases).",
        "Metric: Dice per test case against the manual label.",
        "Predicted fat masks were then cleaned with the HU window -190 to -30 and the",
        "   atrial cavity was subtracted, so only real fat outside the atrium is measured.",
    ]
    segmentation += (
        [f"{name}: Dice {mean:.3f} +- {sd:.3f} on {n} test cases" for name, mean, sd, n in bars]
        if bars
        else ["DINOv2 + LoRA foundation model: run in progress, Dice to be added."]
    )
    add_bullets(presentation, "Segmentation and how it was evaluated", segmentation)
    if dice_image:
        add_picture(presentation, "Segmentation Dice on the test set", dice_image)

    if table is not None:
        groups = [g for g in ["AF", "HF", "stroke"] if (table["group"] == g).any()]
        image = figure_measurements(table, groups, figures / "measurements.png")
        tests = group_tests(table, groups)
        if image:
            add_picture(
                presentation,
                "Plain measurements per group",
                image,
                "Boxes without outliers; control group on the left of each panel.",
            )
        if not tests.empty:
            add_table(
                presentation,
                "Measurements against the controls (Mann-Whitney)",
                tests.sort_values("p_value").head(12).round(3),
            )
        scores = tiers(table, groups)
        if not scores.empty:
            image = figure_auc(scores, figures / "auc.png")
            if image:
                add_picture(
                    presentation,
                    "What predicts disease: measurement, index or radiomics",
                    image,
                    "Logistic regression, 5-fold stratified cross-validation.",
                )
            best = (
                scores.sort_values("auc", ascending=False)
                .groupby("group")
                .head(3)
                .sort_values(["group", "auc"], ascending=[True, False])
            )
            add_table(presentation, "Best model per group", best.round(3))
            scores.to_excel(args.output / "pptx_auc_tiers.xlsx", index=False)

    add_bullets(
        presentation,
        "Where this stands",
        [
            "Plain volumes and mean attenuation are weak on their own.",
            "Composite indices of EAT, LA and LA-EAT volume add some separation.",
            "LA first order radiomics with a texture feature gives the strongest",
            "   imaging-only models, but the test groups are small, so the AUCs are",
            "   cross-validated estimates and need an external cohort.",
            "Next: DINOv2 + LoRA segmentation Dice against nnU-Net on the same 10 test",
            "   cases, then the same radiomics on the new masks, then the large",
            "   non-contrast cohort.",
        ],
    )

    args.output.mkdir(parents=True, exist_ok=True)
    path = args.output / "SCAPIS_CT_summary.pptx"
    presentation.save(path)
    print(f"PowerPoint : {path}")
    print(f"Figures    : {figures}")


if __name__ == "__main__":
    main()
